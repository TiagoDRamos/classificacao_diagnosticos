import argparse
import csv
import os
import re
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl
from requests import RequestException

from icd11 import ICD11


PROJECT_ROOT = Path(__file__).resolve().parent.parent


def resolver_caminho_saida(caminho):
    caminho = Path(caminho)
    if not caminho.is_absolute() and caminho.parent == Path("."):
        caminho = PROJECT_ROOT / "outputs" / caminho
    caminho.parent.mkdir(parents=True, exist_ok=True)
    return caminho


def normalizar(valor):
    """Devolve uma chave consistente para comparar valores do Excel."""
    if valor is None:
        return None

    texto = re.sub(r"\s+", " ", str(valor).strip())
    return texto.casefold() or None


def normalizar_cabecalho(valor):
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = texto.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", texto).strip().casefold()


def eh_coluna_secundaria(nome):
    return bool(
        re.fullmatch(
            r"diagnostico secundario(?: \d+)?",
            normalizar_cabecalho(nome),
        )
    )


ALIASES_DIAGNOSTICOS = {
    "catarata": "Catarata",
    "cataratas": "Catarata",
    "cotovelo de tenista": "Cotovelo de tenista",
    "gastrite h, pilory +": "Gastrite H. pylori +",
    "gastrite h.pylori +": "Gastrite H. pylori +",
    "hemorroidas": "Hemorroidas",
    "hemorróidas": "Hemorroidas",
    "herpes zooster": "Herpes zoster",
    "herpes zoster": "Herpes zoster",
    "hiperuricemia": "Hiperuricemia",
    "hiperuricémia": "Hiperuricemia",
    "hipotiroidismo": "Hipotiroidismo Primário",
    "hipotiroidismo primário": "Hipotiroidismo Primário",
    "hénia inguinal": "Hérnia inguinal",
    "hérnia inguinal": "Hérnia inguinal",
    "insuficiência venosa periférica": "Insuficiência Venosa Periférica",
    "insuficiência venosa perifética": "Insuficiência Venosa Periférica",
    "angiodisplasia do cólon": "Angiodisplasia do Cólon",
    "fibromioma do útero": "Fibromioma do útero",
    "fibromioma utero": "Fibromioma do útero",
    "leiomioma útero": "Leiomioma do útero",
    "leiomioma uterino": "Leiomioma do útero",
    "mioma uterino": "Leiomioma do útero",
    "pti": "Púrpura trombocitopénica imune",
    "púrpura trombocitopénica imune": "Púrpura trombocitopénica imune",
    "polineuropatia não especificada": "Polineuropatia axonal sensitiva",
    "osteoartrose da anca": "Osteoartrose da anca",
    "osteoatrose anca": "Osteoartrose da anca",
    "osteoartrose do joelho": "Osteoartrose do joelho",
    "osteoartrose joelho": "Osteoartrose do joelho",
    "osteoartrose joelhos": "Osteoartrose do joelho",
    "osteoatrose do joelho": "Osteoartrose do joelho",
    "perturbação ansiedade": "Perturbação de ansiedade",
    "perturbação de ansiedade": "Perturbação de ansiedade",
    "perurbação de ansiedade": "Perturbação de ansiedade",
    "perturbação do sono": "Perturbação do sono",
    "perurbação do sono": "Perturbação do sono",
    "síndrome de rendu-osler-weber": "Síndrome de Rendu-Osler-Weber",
    "síndrome osler weber rendu": "Síndrome de Rendu-Osler-Weber",
    "síndrome do canal cárpico": "Síndrome do canal cárpico",
    "síndrome do tunel cárpico": "Síndrome do canal cárpico",
    "síndrome vertiginosa": "Síndrome Vertiginoso",
    "síndrome vertiginoso": "Síndrome Vertiginoso",
    "fa 21 artrite psoriática": "Artrite Psoriática",
    "cboz doença do interstício pulmonar": "Doença do Interstício Pulmonar",
    "osteoartose do joelho": "Osteoartrose do joelho",
}

PADRAO_CODIGO_ICD11 = re.compile(
    r"^\s*([0-9A-Z][0-9A-Z]{1,5}(?:\.[0-9A-Z]+)?\.?)\s+(.+?)\s*$",
    re.IGNORECASE,
)

CODIGOS_ICD11_FALLBACK = {
    "anemia em estudo": "3A9Z",
    "anemia ferropénica": "3A00",
    "doença do interstício pulmonar": "CB0Z",
    "celulite de repetição": "1B70",
    "doença hepática gordurasa soe": "DB92.Z",
    "doença pulmonar intesticial": "CB0Z",
    "espasmo esofágico difuso": "DA21.2",
    "ferropenia": "5B5K.0",
    "hiperferritinémia": "5C64.1Z",
    "hipotiroidismo": "5A00",
    "hipotiroidismo primário": "5A00",
    "linfadenite granulomatosa": "BD90",
    "nódulo da suprarrenal": "5A76.Y",
    "perda ponderal": "MG43.5",
    "polineuropatia axonal sensitiva": "8C0Z",
    "síndrome de rendu-osler-weber": "LA90.00",
    "artrite psoriática": "FA21",
    "fa 21 artrite psoriática": "FA21",
    "mgus": "2A83.0",
    "hiperolasia benigna da próstata":"GA90",
    "adenocarcinoma de brônqio ou pulmão":"2C25.0"
}


def canonicalizar_diagnostico(valor):
    if valor is None:
        return None

    texto = re.sub(r"\s+", " ", str(valor).strip()).rstrip(";").strip()
    if not texto:
        return None

    return ALIASES_DIAGNOSTICOS.get(normalizar(texto), texto)


def extrair_codigo_icd11(diagnostico):
    correspondencia = PADRAO_CODIGO_ICD11.match(diagnostico or "")
    if not correspondencia:
        return None

    codigo = correspondencia.group(1).upper().rstrip(".")
    if not any(carater.isdigit() for carater in codigo):
        return None

    titulo = correspondencia.group(2).strip()
    return codigo, titulo


def remover_codigo_icd11(diagnostico):
    codigo = extrair_codigo_icd11(diagnostico)
    if codigo is None:
        return diagnostico

    return codigo[1]


def grupo_a_partir_do_codigo(codigo):
    primeiro_carater = codigo[0].upper()
    if primeiro_carater.isdigit():
        return f"{int(primeiro_carater):02d}"

    grupos_icd11 = {
        letra: grupo
        for grupo, letra in enumerate("ABCDEFGHJKLMNPQRSTUVWXYZ", start=10)
    }
    try:
        return f"{grupos_icd11[primeiro_carater]:02d}"
    except KeyError as error:
        raise ValueError(f"Código ICD-11 inválido: {codigo}") from error


def classificacao_local(diagnostico, codigo):
    return {
        "match": {
            "chapter": grupo_a_partir_do_codigo(codigo),
            "code": codigo,
            "title": diagnostico,
        },
        "confidence": "local",
    }


def encontrar_coluna(cabecalhos, nome):
    nome_normalizado = normalizar(nome)
    for indice, cabecalho in enumerate(cabecalhos):
        if normalizar(cabecalho) == nome_normalizado:
            return indice
    raise ValueError(f"Coluna '{nome}' não encontrada na primeira folha.")


def obter_sheet_id(url):
    correspondencia = re.search(
        r"docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)",
        str(url),
    )
    if not correspondencia:
        return None
    return correspondencia.group(1)


def ler_google_sheet(sheet_id):
    credenciais = os.environ.get("GOOGLE_APPLICATION_CREDENTIALS")
    if not credenciais:
        raise ValueError(
            "Para ler uma folha privada, define a variável "
            "GOOGLE_APPLICATION_CREDENTIALS com o caminho do JSON da conta de serviço."
        )

    try:
        import gspread
    except ImportError as error:
        raise RuntimeError(
            "Falta a dependência gspread. Instala as dependências com "
            "'python -m pip install -r requirements.txt'."
        ) from error

    cliente = gspread.service_account(
        filename=credenciais,
        scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"],
    )
    return cliente.open_by_key(sheet_id).sheet1.get_all_values()


def preparar_contagem(ficheiro_entrada):
    sheet_id = obter_sheet_id(ficheiro_entrada)
    if sheet_id:
        linhas_sheet = ler_google_sheet(sheet_id)
        cabecalhos = linhas_sheet[0] if linhas_sheet else None
        rows = linhas_sheet[1:]
    else:
        workbook = openpyxl.load_workbook(
            ficheiro_entrada,
            read_only=True,
            data_only=True,
        )
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        cabecalhos = next(rows, None)

    if cabecalhos is None:
        raise ValueError("A primeira folha está vazia.")

    indice_principal = encontrar_coluna(cabecalhos, "Diagnóstico principal")
    indices_secundarios = [
        indice
        for indice, cabecalho in enumerate(cabecalhos)
        if eh_coluna_secundaria(cabecalho)
    ]

    contagens_principais = Counter()
    contagens_secundarias = Counter()
    nomes_diagnosticos = {}

    for linha in rows:
        valor_principal = (
            linha[indice_principal]
            if indice_principal < len(linha)
            else None
        )
        valor_principal = canonicalizar_diagnostico(valor_principal)
        chave_principal = normalizar(valor_principal)

        if chave_principal is not None:
            contagens_principais[chave_principal] += 1
            nomes_diagnosticos.setdefault(
                chave_principal,
                str(valor_principal).strip(),
            )

        for indice_secundario in indices_secundarios:
            if indice_secundario >= len(linha):
                continue

            valor_secundario = canonicalizar_diagnostico(linha[indice_secundario])
            chave_secundaria = normalizar(valor_secundario)

            if chave_secundaria is not None:
                contagens_secundarias[chave_secundaria] += 1
                nomes_diagnosticos.setdefault(
                    chave_secundaria,
                    str(valor_secundario).strip(),
                )

    cabecalhos_resultado = [
        "Diagnóstico",
        "Total em Diagnóstico principal",
        "Total em Diagnóstico secundário",
    ]
    dados_resultado = []

    todas_as_chaves = set(contagens_principais) | set(contagens_secundarias)

    for chave in sorted(
        todas_as_chaves,
        key=lambda item: nomes_diagnosticos[item].casefold(),
    ):
        dados_resultado.append([
            nomes_diagnosticos[chave],
            contagens_principais[chave],
            contagens_secundarias[chave],
        ])

    return cabecalhos_resultado, dados_resultado


def guardar_excel(cabecalhos, dados, ficheiro_saida):
    Path(ficheiro_saida).parent.mkdir(parents=True, exist_ok=True)
    resultado = openpyxl.Workbook()
    resultado_sheet = resultado.active
    resultado_sheet.title = "Contagem"
    resultado_sheet.append(cabecalhos)
    for linha in dados:
        resultado_sheet.append(linha)

    resultado_sheet.freeze_panes = "A2"
    resultado_sheet.auto_filter.ref = resultado_sheet.dimensions
    resultado_sheet.column_dimensions["A"].width = 35
    resultado_sheet.column_dimensions["B"].width = 30
    resultado_sheet.column_dimensions["C"].width = 28
    resultado.save(ficheiro_saida)
    return Path(ficheiro_saida)


def guardar_csv_com_icd11(dados, ficheiro_saida):
    icd = ICD11()
    erro_icd11 = None
    cabecalhos = [
        "Diagnóstico",
        "Grupo ICD-11",
        "Código ICD-11",
        "Correspondência ICD-11",
        "Confiança",
        "Total em Diagnóstico principal",
        "Total em Diagnóstico secundário",
    ]
    linhas = []
    dados_agregados = {}

    for diagnostico, total_principal, total_secundario in dados:
        nome = canonicalizar_diagnostico(remover_codigo_icd11(diagnostico))
        chave = normalizar(nome)
        if chave not in dados_agregados:
            dados_agregados[chave] = [nome, diagnostico, 0, 0]
        dados_agregados[chave][2] += total_principal
        dados_agregados[chave][3] += total_secundario

    for nome, diagnostico_original, total_principal, total_secundario in sorted(
        dados_agregados.values(),
        key=lambda linha: linha[0].casefold(),
    ):
        codigo_local = extrair_codigo_icd11(diagnostico_original)
        codigo_fallback = CODIGOS_ICD11_FALLBACK.get(normalizar(nome))

        if codigo_local:
            codigo, titulo = codigo_local
            classificacao = classificacao_local(titulo, codigo)
        elif codigo_fallback:
            classificacao = classificacao_local(nome, codigo_fallback)
        else:
            classificacao = {"match": None, "confidence": "none"}

            if erro_icd11 is None:
                try:
                    classificacao = icd.classify(nome)
                except RequestException as error:
                    erro_icd11 = error
                    classificacao["confidence"] = "indisponível"

            if erro_icd11 is not None:
                classificacao = {
                    "match": None,
                    "confidence": "indisponível",
                }

        match = classificacao.get("match")

        if match is None:
            grupo = "Serviço ICD-11 indisponível" if erro_icd11 else "Não encontrado"
            codigo = ""
            correspondencia = ""
        else:
            grupo = match.get("chapter") or "Não indicado"
            codigo = match.get("code") or ""
            correspondencia = match.get("title") or ""

        linhas.append([
            nome,
            grupo,
            codigo,
            correspondencia,
            classificacao.get("confidence", "none"),
            total_principal,
            total_secundario,
        ])

    return guardar_csv(cabecalhos, linhas, ficheiro_saida)


def guardar_csv(cabecalhos, dados, ficheiro_saida):
    Path(ficheiro_saida).parent.mkdir(parents=True, exist_ok=True)
    with open(ficheiro_saida, "w", newline="", encoding="utf-8-sig") as ficheiro:
        escritor = csv.writer(ficheiro, delimiter=";")
        escritor.writerow(cabecalhos)
        escritor.writerows(dados)
    return Path(ficheiro_saida)


def main():
    parser = argparse.ArgumentParser(
        description="Conta diagnósticos principais e as suas ocorrências nas restantes colunas."
    )
    parser.add_argument(
        "ficheiro_entrada",
        help="Caminho para um Excel local ou URL de uma folha Google Sheets",
    )
    parser.add_argument(
        "-o",
        "--saida",
        default="diagnosticos_contagem.xlsx",
        help="Nome do ficheiro de saída (por defeito, guardado em outputs/)",
    )
    parser.add_argument(
        "--csv",
        action="store_true",
        help="Exporta para CSV em vez de Excel",
    )
    parser.add_argument(
        "--icd11-saida",
        metavar="FICHEIRO",
        help="Nome do CSV ICD-11 (por defeito, guardado em outputs/)",
    )
    args = parser.parse_args()

    cabecalhos, dados = preparar_contagem(args.ficheiro_entrada)
    ficheiro_saida = resolver_caminho_saida(args.saida)
    if args.csv:
        ficheiro_saida = guardar_csv(cabecalhos, dados, ficheiro_saida)
    else:
        ficheiro_saida = guardar_excel(cabecalhos, dados, ficheiro_saida)
    print(f"Resultado guardado em: {ficheiro_saida}")

    if args.icd11_saida:
        caminho_icd11 = resolver_caminho_saida(args.icd11_saida)
        ficheiro_icd11 = guardar_csv_com_icd11(dados, caminho_icd11)
        print(f"Resultado ICD-11 guardado em: {ficheiro_icd11}")


if __name__ == "__main__":
    main()